import os
import shutil
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, BackgroundTasks
from pydantic import BaseModel
from sqlalchemy.orm import Session

from core.database import get_db
from core.dependencies import get_current_admin
from models.models.knowledge_base import KBSection, KBItem, KBDocument
from services.embedding_service import embed_and_store, search_kb_rag, ensure_embedding_column

router = APIRouter(prefix="/kb", tags=["Knowledge Base"])

UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "..", "uploads", "kb_docs")
os.makedirs(UPLOAD_DIR, exist_ok=True)


# ── Schemas ───────────────────────────────────────────────────────────────────

class KBItemCreate(BaseModel):
    question: str
    answer: str
    sort_order: int = 0


class KBItemUpdate(BaseModel):
    question: Optional[str] = None
    answer: Optional[str] = None
    is_active: Optional[bool] = None
    sort_order: Optional[int] = None


# ── Sections ──────────────────────────────────────────────────────────────────

@router.get("/sections")
def list_sections(db: Session = Depends(get_db), _=Depends(get_current_admin)):
    sections = (
        db.query(KBSection)
        .filter(KBSection.is_active == True)
        .order_by(KBSection.sort_order)
        .all()
    )
    return [
        {
            "id": s.id,
            "name": s.name,
            "slug": s.slug,
            "description": s.description,
            "sort_order": s.sort_order,
        }
        for s in sections
    ]


# ── Items ─────────────────────────────────────────────────────────────────────

@router.get("/sections/{slug}/items")
def list_items(slug: str, db: Session = Depends(get_db), _=Depends(get_current_admin)):
    section = db.query(KBSection).filter(KBSection.slug == slug).first()
    if not section:
        raise HTTPException(status_code=404, detail="Section not found")

    items = (
        db.query(KBItem)
        .filter(KBItem.section_id == section.id, KBItem.is_active == True)
        .order_by(KBItem.sort_order)
        .all()
    )
    return {
        "section": {"id": section.id, "name": section.name, "slug": section.slug},
        "items": [
            {
                "id": i.id,
                "question": i.question,
                "answer": i.answer,
                "sort_order": i.sort_order,
                "created_at": i.created_at,
            }
            for i in items
        ],
    }


@router.post("/sections/{slug}/items", status_code=201)
def add_item(
    slug: str,
    body: KBItemCreate,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    _=Depends(get_current_admin),
):
    section = db.query(KBSection).filter(KBSection.slug == slug).first()
    if not section:
        raise HTTPException(status_code=404, detail="Section not found")

    item = KBItem(
        section_id=section.id,
        question=body.question,
        answer=body.answer,
        sort_order=body.sort_order,
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    background_tasks.add_task(embed_and_store, item.id, item.question, item.answer, db)
    return {"id": item.id, "message": "Item added"}


@router.put("/items/{item_id}")
def update_item(
    item_id: int,
    body: KBItemUpdate,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    _=Depends(get_current_admin),
):
    item = db.query(KBItem).filter(KBItem.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")

    if body.question is not None:
        item.question = body.question
    if body.answer is not None:
        item.answer = body.answer
    if body.is_active is not None:
        item.is_active = body.is_active
    if body.sort_order is not None:
        item.sort_order = body.sort_order

    db.commit()
    if body.question is not None or body.answer is not None:
        background_tasks.add_task(embed_and_store, item.id, item.question, item.answer, db)
    return {"message": "Updated"}


@router.delete("/items/{item_id}")
def delete_item(
    item_id: int,
    db: Session = Depends(get_db),
    _=Depends(get_current_admin),
):
    item = db.query(KBItem).filter(KBItem.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")

    item.is_active = False
    db.commit()
    return {"message": "Deleted"}


# ── Documents (PDF / DOCX upload) ─────────────────────────────────────────────

@router.post("/sections/{slug}/upload", status_code=201)
def upload_document(
    slug: str,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _=Depends(get_current_admin),
):
    section = db.query(KBSection).filter(KBSection.slug == slug).first()
    if not section:
        raise HTTPException(status_code=404, detail="Section not found")

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in (".pdf", ".docx"):
        raise HTTPException(status_code=400, detail="Only PDF and DOCX files allowed")

    save_path = os.path.join(UPLOAD_DIR, f"{section.slug}_{file.filename}")
    with open(save_path, "wb") as f:
        shutil.copyfileobj(file.file, f)

    extracted = _extract_text(save_path, ext)

    doc = KBDocument(
        section_id=section.id,
        file_name=file.filename,
        file_path=save_path,
        file_type=ext.lstrip("."),
        extracted_text=extracted,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return {
        "id": doc.id,
        "file_name": doc.file_name,
        "file_type": doc.file_type,
        "extracted_length": len(extracted) if extracted else 0,
        "message": "Uploaded and text extracted",
    }


@router.get("/sections/{slug}/documents")
def list_documents(slug: str, db: Session = Depends(get_db), _=Depends(get_current_admin)):
    section = db.query(KBSection).filter(KBSection.slug == slug).first()
    if not section:
        raise HTTPException(status_code=404, detail="Section not found")

    docs = (
        db.query(KBDocument)
        .filter(KBDocument.section_id == section.id, KBDocument.is_active == True)
        .order_by(KBDocument.uploaded_at.desc())
        .all()
    )
    return [
        {
            "id": d.id,
            "file_name": d.file_name,
            "file_type": d.file_type,
            "uploaded_at": d.uploaded_at,
            "has_text": bool(d.extracted_text),
        }
        for d in docs
    ]


@router.delete("/documents/{doc_id}")
def delete_document(
    doc_id: int,
    db: Session = Depends(get_db),
    _=Depends(get_current_admin),
):
    doc = db.query(KBDocument).filter(KBDocument.id == doc_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    doc.is_active = False
    db.commit()
    return {"message": "Deleted"}


# ── RAG — reindex all items ───────────────────────────────────────────────────

@router.post("/reindex")
def reindex_kb(db: Session = Depends(get_db), _=Depends(get_current_admin)):
    """Embed every active KB item. Run once after setup or after bulk edits."""
    ensure_embedding_column(db)
    items = db.query(KBItem).filter(KBItem.is_active == True).all()
    ok = 0
    for item in items:
        if embed_and_store(item.id, item.question, item.answer, db):
            ok += 1
    return {"total": len(items), "embedded": ok, "message": "Reindex complete"}


@router.get("/search")
def search_kb(query: str, top_k: int = 6, db: Session = Depends(get_db), _=Depends(get_current_admin)):
    """Semantic search across all KB items."""
    results = search_kb_rag(query, db, top_k=top_k)
    return {"query": query, "results": results}


# ── Text extraction helpers ───────────────────────────────────────────────────

def _extract_text(path: str, ext: str) -> str:
    try:
        if ext == ".pdf":
            import pdfplumber
            text_parts = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        text_parts.append(t)
            return "\n".join(text_parts)

        elif ext == ".docx":
            from docx import Document
            doc = Document(path)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception:
        pass
    return ""
